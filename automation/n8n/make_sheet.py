#!/usr/bin/env python3
"""Build the final Google-Sheet column template for the Watad Unified Engine."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WF = "/home/user/hc/automation/n8n/Watad_Unified_Engine_Brand_KB_driven.json"
OUT = "/home/user/hc/automation/n8n/Watad_Engine_Sheet_Columns.xlsx"

NAVY = "0F1D2D"
TEAL = "00A38C"
GOLD = "C9A84C"
LIGHT = "F4F7FA"
GREY = "8A98A6"

F = "Arial"
thin = Side(style="thin", color="D8DEE5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------------------------------------------------------------- definitions
# owner: engine | human ; new: added by this round of work
COLS = {
    "Brand_KB": [
        ("brand", "human", False, "المُعرّف المختصر للبراند بحروف صغيرة — المفتاح اللي كل حاجة بتتربط بيه"),
        ("brand_name", "human", False, "الاسم المعروض للبراند زي ما يظهر في المحتوى"),
        ("active", "human", False, "yes / no — الصف بيتوقف بـ no من غير حذف"),
        ("zernio_profile_id", "human", False, "أقوى إشارة للمطابقة. أكتر من ID تتفصل بفاصلة"),
        ("zernio_profile_names", "human", False, "أسماء البروفايل في Zernio، بفاصلة. مطابقة تقريبية"),
        ("zernio_account_ids", "human", False, "IDs الحسابات، بفاصلة"),
        ("kb", "human", False, "قاعدة المعرفة — دي مصدر الحقيقة الوحيد لبوابة الامتثال"),
        ("services", "human", False, "الخدمات، بفاصلة أو فاصلة منقوطة"),
        ("faq", "human", False, "أسئلة شائعة وإجاباتها المعتمدة"),
        ("never_say", "human", False, "ممنوع قوله — أي مصطلح هنا بيحظر المحتوى قبل توليد الميديا"),
        ("tone", "human", False, "نبرة البراند"),
        ("audience", "human", False, "الجمهور المستهدف"),
        ("focus", "human", False, "مجالات التركيز"),
        ("anti_hype", "human", False, "الأسلوب المرفوض صراحةً"),
        ("visual_style", "human", False, "الهوية البصرية — بتغذّي photo_prompt"),
        ("hashtags", "human", False, "بنك الهاشتاجات المعتمد"),
        ("escalation_email", "human", False, "كل إشعارات الموافقة والتصعيد بتروح هنا"),
        ("allow_auto_reply", "human", False, "yes / no — no بتوقف الرد الآلي لهذا البراند"),
    ],
    "Content": [
        ("run_id", "engine", False, "مفتاح الصف. كل العقد بتطابق عليه"),
        ("brand", "engine", False, "البراند اللي الـ Registry حلّه"),
        ("profileId", "engine", False, "بروفايل Zernio — بيحدد حسابات النشر"),
        ("profileName", "engine", False, "اسم البروفايل — بديل الـ ID في المطابقة"),
        ("approval_status", "human", False, "PENDING / approved / rejected — أزرار الإيميل بتكتب هنا"),
        ("publish_status", "engine", False, "Waiting / PUBLISHING / Scheduled / FAILED / NO_ACCOUNT"),
        ("compliance_status", "engine", True, "PASS / BLOCKED — نتيجة بوابة الامتثال"),
        ("compliance_notes", "engine", True, "سبب الحظر + نص المخالفة حرفياً"),
        ("approval_token", "engine", True, "توكن أزرار الموافقة. متلمسوش"),
        ("approved_at", "engine", True, "وقت الضغط على موافق/ارفض"),
        ("topic_title", "engine", False, "الموضوع. ده اللي ذاكرة المواضيع بتقراه"),
        ("facebook_post", "engine", False, "نسخة فيسبوك"),
        ("instagram_post", "engine", False, "نسخة إنستجرام"),
        ("linkedin_post", "engine", False, "نسخة لينكدإن"),
        ("twitter_post", "engine", False, "نسخة تويتر — ٢٨٠ حرف حد أقصى"),
        ("tiktok_caption", "engine", False, "كابشن تيك توك — لازم فيديو معاه"),
        ("photo_prompt", "engine", False, "برومبت الصورة"),
        ("video_prompt", "engine", False, "برومبت الفيديو — بيتولّد بعد الموافقة بس"),
        ("warnings", "engine", False, "ملاحظات المحلّل على مخرجات النموذج"),
        ("Photo_Link", "engine", False, "الصورة الأساسية على Zernio"),
        ("Photo_Links", "engine", False, "كل شرائح الكاروسيل بالترتيب، بفاصلة"),
        ("Video_Link", "engine", False, "الفيديو على Zernio"),
        ("Drive_Photo_Link", "engine", True, "نسخة الصورة الأرشيفية على Drive"),
        ("Drive_Video_Link", "engine", True, "نسخة الفيديو الأرشيفية على Drive"),
        ("render_url", "engine", False, "لينك hcti المؤقت لصفوف الديزاين"),
        ("design_type", "engine", False, "yehia / navid / playground"),
        ("output_type", "engine", False, "image / carousel"),
        ("zernio_post_id", "engine", False, "مُعرّف البوست بعد الجدولة — مفتاح سحب الأداء"),
        ("scheduled_for", "engine", False, "وقت النشر بتوقيت الرياض"),
        ("published_platforms", "engine", False, "المنصات اللي اتنشر عليها فعلاً"),
        ("publish_error", "engine", False, "سبب فشل النشر"),
        ("perf_stage", "engine", True, "24h / 72h — مؤشّر سحب الأداء، بيمنع التكرار"),
        ("perf_last_pull", "engine", True, "وقت آخر سحب"),
        ("perf_engagement", "engine", True, "التفاعل من آخر لقطة"),
        ("perf_impressions", "engine", True, "الظهور من آخر لقطة"),
    ],
    "Leads": [
        ("sender_id", "engine", False, "مفتاح الصف — مُعرّف المرسل من المنصة"),
        ("platform", "engine", False, "المنصة"),
        ("full_name", "engine", False, "الاسم"),
        ("email", "engine", False, "الإيميل — فاضي أفضل من مخترع"),
        ("phone", "engine", False, "التليفون"),
        ("title", "engine", False, "المسمى الوظيفي"),
        ("company", "engine", False, "الشركة"),
        ("linkedin_url", "engine", False, "بروفايل لينكدإن"),
        ("city", "engine", False, "المدينة"),
        ("country", "engine", False, "الدولة"),
        ("seniority", "engine", False, "المستوى الوظيفي"),
        ("first_seen", "engine", False, "أول ظهور"),
        ("last_messages", "engine", False, "سجل المحادثة — محرك الردود بيقرا آخر ٣٠٠٠ حرف"),
        ("source", "engine", True, "المصدر. كان مكتوب غلط sourse — لازم تغيّر اسم العمود"),
        ("photo", "engine", False, "لينك صورة البروفايل على Drive"),
        ("url", "engine", False, "لينك البروفايل"),
        ("osint_findings", "engine", False, "خلاصة بحث الـ OSINT"),
    ],
    "Reply_Approval": [
        ("run_id", "engine", False, "مفتاح الصف"),
        ("status", "human", False, "PENDING / approved / SENDING / SENT / rejected / REJECTED / FAILED / INVALID"),
        ("approval_token", "engine", True, "توكن أزرار الموافقة. متلمسوش"),
        ("approved_at", "engine", True, "وقت الضغط على موافق/ارفض"),
        ("timestamp", "engine", False, "وقت وصول الرسالة"),
        ("brand", "engine", False, "البراند"),
        ("platform", "engine", False, "المنصة"),
        ("event_type", "engine", False, "message.received / comment.received"),
        ("sender_id", "engine", False, "مُعرّف المرسل"),
        ("sender_name", "engine", False, "اسم المرسل"),
        ("sender_company", "engine", False, "شركة المرسل"),
        ("sender_title", "engine", False, "منصب المرسل"),
        ("original_message", "engine", False, "الرسالة الأصلية"),
        ("ai_draft_reply", "engine", False, "المسودة. بتعدّي على فحص أمان تاني قبل الإرسال"),
        ("reason", "engine", False, "سبب التوجيه للمراجعة"),
        ("conversation_key", "engine", False, "مفتاح المحادثة"),
        ("conversation_id", "engine", False, "لازم للـ DM"),
        ("comment_id", "engine", False, "لازم للكومنت"),
        ("post_id", "engine", False, "لازم للكومنت"),
        ("account_id", "engine", False, "حساب Zernio اللي هيبعت"),
        ("request_id", "engine", False, "مفتاح منع التكرار"),
        ("send_error", "engine", False, "سبب الفشل أو الرفض"),
        ("sent_at", "engine", False, "وقت الإرسال"),
    ],
    "Events_Log": [
        ("event_key", "engine", False, "مفتاح منع التكرار — Zernio بتعيد الإرسال"),
        ("run_id", "engine", False, "مُعرّف التشغيلة"),
        ("timestamp", "engine", False, "ISO. التدوير الشهري بيعتمد عليه"),
        ("brand", "engine", False, "البراند"),
        ("platform", "engine", False, "المنصة"),
        ("event_type", "engine", False, "نوع الحدث"),
        ("sender_id", "engine", False, "مُعرّف المرسل"),
        ("sender_name", "engine", False, "اسم المرسل"),
        ("category", "engine", False, "reply / approval / handoff / abuse"),
        ("intent", "engine", False, "نية الرسالة"),
        ("sentiment", "engine", False, "positive / neutral / negative"),
        ("lead_quality", "engine", False, "hot / warm / cold"),
        ("safe_to_send", "engine", False, "نتيجة بوابة الأمان"),
        ("safety_reason", "engine", False, "سبب المنع"),
        ("message", "engine", False, "نص الرسالة"),
        ("reply_sent", "engine", False, "الرد اللي اتبعت"),
        ("target", "engine", False, "الوجهة"),
        ("http_status", "engine", False, "كود استجابة Zernio"),
        ("route_reason", "engine", False, "سبب قرار الراوتر"),
        ("needs_human", "engine", False, "محتاج تدخل بشري"),
        ("answerable", "engine", False, "مغطّى بقاعدة المعرفة"),
        ("asked_for_human", "engine", False, "طلب التحدث لإنسان صراحةً"),
    ],
    "Handoff": [
        ("run_id", "engine", False, "مُعرّف التشغيلة"),
        ("status", "engine", False, "حالة التسليم"),
        ("timestamp", "engine", False, "الوقت"),
        ("brand", "engine", False, "البراند"),
        ("platform", "engine", False, "المنصة"),
        ("event_type", "engine", False, "نوع الحدث"),
        ("sender_id", "engine", False, "مُعرّف المرسل"),
        ("sender_name", "engine", False, "اسم المرسل"),
        ("message", "engine", False, "الرسالة"),
        ("intent", "engine", False, "النية"),
        ("sentiment", "engine", False, "المشاعر"),
        ("lead_quality", "engine", False, "جودة الـ lead"),
        ("urgency", "engine", False, "الإلحاح"),
        ("reply_draft", "engine", False, "مسودة الرد"),
        ("reason", "engine", False, "سبب التصعيد"),
        ("conversation_id", "engine", False, "مُعرّف المحادثة"),
        ("comment_id", "engine", False, "مُعرّف الكومنت"),
        ("post_id", "engine", False, "مُعرّف البوست"),
        ("account_id", "engine", False, "مُعرّف الحساب"),
    ],
    "Performance": [
        ("run_id", "engine", True, "الصف في تبويب Content"),
        ("brand", "engine", True, "البراند"),
        ("zernio_post_id", "engine", True, "مُعرّف البوست في Zernio"),
        ("stage", "engine", True, "24h أو 72h"),
        ("pulled_at", "engine", True, "وقت السحب"),
        ("topic_title", "engine", True, "الموضوع — عشان تقارن الأداء بالمواضيع"),
        ("platforms", "engine", True, "المنصات"),
        ("impressions", "engine", True, "الظهور"),
        ("reach", "engine", True, "الوصول"),
        ("likes", "engine", True, "الإعجابات"),
        ("comments", "engine", True, "التعليقات"),
        ("shares", "engine", True, "المشاركات"),
        ("saves", "engine", True, "الحفظ"),
        ("clicks", "engine", True, "النقرات"),
        ("views", "engine", True, "المشاهدات"),
        ("follows", "engine", True, "متابعات جديدة"),
        ("engagement", "engine", True, "التفاعل — محسوب لو Zernio ما رجّعوش"),
        ("raw", "engine", True, "أول ٩٠٠ حرف من رد Zernio للتشخيص"),
    ],
}
COLS["Events_Log_Archive"] = [(c, o, True, desc) for c, o, _, desc in COLS["Events_Log"]]

EXAMPLES = {
    "Brand_KB": ["watad", "Watad Digital", "yes", "68f1a2b3c4d5e6f708192a3b",
                 "Watad, Watad Digital", "acc_li_001, acc_fb_001",
                 "شركة سعودية للتقنية المؤسسية تأسست 2013. المقر الخبر ودبي. +966138870349",
                 "أمن OT/ICS; الأتمتة الصناعية; المدن الذكية; الذكاء الاصطناعي",
                 "س: مواعيد العمل؟ ج: الأحد–الخميس 8–5 بتوقيت السعودية",
                 "الأسعار, مواعيد التسليم, أسماء العملاء", "تنفيذية، مؤسسية، بمقياس وطني",
                 "مدراء أمن المعلومات والجهات الحكومية", "البنية التحتية الحرجة والمدن الذكية",
                 "ممنوع لغة الثورة الرقمية؛ دايماً الجاهزية التشغيلية",
                 "مؤسسي داكن، كحلي وتركوازي #00C9B1، غرف تحكم",
                 "#CriticalInfrastructure #OTSecurity #SaudiTech",
                 "ops@wataddigital.com", "yes"],
    "Content": ["RUN-watad-8241-0", "watad", "68f1a2b3c4d5e6f708192a3b", "Watad Digital",
                "PENDING", "Waiting", "PASS", "", "k3n8vq2xr7m1pd4wzt6ybs0a",
                "", "لماذا تفشل مشاريع أمن OT قبل التشغيل",
                "نص فيسبوك…", "نص إنستجرام…", "نص لينكدإن…", "نص تويتر…", "كابشن تيك توك…",
                "cinematic dark enterprise control room, 4:5, no text",
                "12s vertical, 3 scenes, no text", "",
                "https://cdn.zernio.com/media/RUN-watad-8241-0.png", "", "",
                "https://drive.google.com/file/d/1AbC.../view", "", "", "", "",
                "", "", "", "", "", "", "", ""],
    "Leads": ["17841400000000000", "instagram", "أحمد الغامدي", "a.ghamdi@example.com",
              "+966500000000", "مدير أمن المعلومات", "شركة المثال",
              "https://linkedin.com/in/example", "الخبر", "السعودية", "director",
              "2026-07-14T09:12:00Z", "سؤال عن تغطية OT للمصانع…", "instagram_dm",
              "https://drive.google.com/file/d/1Xy.../view",
              "https://instagram.com/example", "خبرة 12 سنة في أمن الشبكات الصناعية"],
    "Reply_Approval": ["RPL-msg8f2a91-8241", "PENDING", "q7w2e9r4t1y6u3i8o5p0a2s4", "",
                       "2026-08-02T10:04:11Z", "watad", "instagram", "message.received",
                       "17841400000000000", "أحمد الغامدي", "شركة المثال", "مدير أمن المعلومات",
                       "عايز أعرف تغطيتكم للمصانع", "شكراً لتواصلك…", "غير مغطّى بقاعدة المعرفة",
                       "instagram:conv_88123", "conv_88123", "", "", "acc_ig_001",
                       "3f2a-91bd-4c7e", "", ""],
    "Events_Log": ["ig:msg:8f2a91", "RPL-msg8f2a91-8241", "2026-08-02T10:04:11Z", "watad",
                   "instagram", "message.received", "17841400000000000", "أحمد الغامدي",
                   "reply", "sales_inquiry", "neutral", "warm", "TRUE", "",
                   "عايز أعرف تغطيتكم للمصانع", "شكراً لتواصلك…", "dm", "200",
                   "مغطّى بالكامل بقاعدة المعرفة", "FALSE", "TRUE", "FALSE"],
    "Handoff": ["RPL-cmt41c8b-8239", "OPEN", "2026-08-02T09:40:02Z", "navid", "linkedin",
                "comment.received", "ACoAAB1x", "سارة القحطاني", "ممكن أكلم حد من المبيعات؟",
                "sales_inquiry", "positive", "hot", "high", "سيتواصل معك زميل…",
                "طلب التحدث لشخص صراحةً", "", "7241889302", "7241889301", "acc_li_001"],
    "Performance": ["RUN-watad-8241-0", "watad", "68f1a2b3c4d5e6f7", "24h",
                    "2026-08-03T12:00:00Z", "لماذا تفشل مشاريع أمن OT قبل التشغيل",
                    "linkedin,facebook", 1240, 890, 37, 5, 3, 2, 18, 0, 1, 47,
                    '{"impressions":1240,"reach":890}'],
}
EXAMPLES["Events_Log_Archive"] = EXAMPLES["Events_Log"]

# ------------------------------------------------------------------ integrity
wf = json.load(open(WF))
written = {}
for n in wf["nodes"]:
    if n["type"] != "n8n-nodes-base.googleSheets":
        continue
    p = n["parameters"]
    sn = p.get("sheetName") or {}
    name = sn.get("cachedResultName") or sn.get("value")
    name = {"gid=0": "Leads", 1308504207: "Content"}.get(name, str(name))
    for k in (p.get("columns", {}).get("value") or {}):
        written.setdefault(name, set()).add(k)
    for f in (p.get("filtersUI", {}).get("values") or []):
        written.setdefault(name, set()).add(f["lookupColumn"])

for tab, cols in COLS.items():
    declared = {c for c, *_ in cols}
    used = written.get(tab, set())
    missing = used - declared
    assert not missing, f"{tab}: workflow uses columns missing from the template: {missing}"
print("integrity: every column the workflow touches is in the template")

# ---------------------------------------------------------------------- build
wb = Workbook()
wb.remove(wb.active)


def style_header(ws, cols):
    for i, (name, owner, is_new, _desc) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        fill = TEAL if is_new else (GOLD if owner == "human" else NAVY)
        c.font = Font(name=F, bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(30, len(name) + 8))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def example_row(ws, values, ncols):
    for i in range(1, ncols + 1):
        v = values[i - 1] if i - 1 < len(values) else ""
        c = ws.cell(row=2, column=i, value=v)
        c.font = Font(name=F, size=9, italic=True, color=GREY)
        c.alignment = Alignment(vertical="top", wrap_text=False)
        c.border = BORDER
    ws.row_dimensions[2].height = 18


# ---- 1. legend ----
ws = wb.create_sheet("اقرأ-أولاً")
ws.sheet_view.rightToLeft = True
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 96

rows = [
    ("Watad Unified Engine — أعمدة الشيت النهائية", "", "title"),
    ("", "", ""),
    ("إزاي تستخدم الملف", "كل تبويب هنا = تبويب في final_engine. الصف الأول هو رؤوس الأعمدة بالظبط — انسخه كما هو. "
     "الصف التاني مثال للشكل المتوقّع، امسحه قبل التشغيل.", "kv"),
    ("", "", ""),
    ("ألوان رؤوس الأعمدة", "", "head"),
    ("تركوازي", "عمود جديد — لازم تضيفه أو تعيد تسميته", "kv"),
    ("ذهبي", "أنت اللي بتكتب فيه بإيدك", "kv"),
    ("كحلي", "المحرك بيكتبه — سيبه", "kv"),
    ("", "", ""),
    ("المطلوب منك بالظبط", "", "head"),
    ("١. تبويبين جداد", "Performance و Events_Log_Archive — انسخ رؤوسهم من التبويبات هنا. "
     "Events_Log_Archive لازم يكون مطابق لـ Events_Log حرفياً.", "kv"),
    ("٢. إعادة تسمية", "في تبويب Leads: العمود sourse ← source", "kv"),
    ("٣. ١٠ أعمدة في Content", "compliance_status · compliance_notes · approval_token · approved_at · "
     "Drive_Photo_Link · Drive_Video_Link · perf_stage · perf_last_pull · perf_engagement · perf_impressions", "kv"),
    ("٤. عمودين في Reply_Approval", "approval_token · approved_at", "kv"),
    ("٥. Brand_KB", "تأكد إن كل براند له صف نشط بكل الأعمدة. كل المحرّكات بقت تسأل نفس الـ Registry، "
     "و never_say بقى بيحظر المحتوى فعلياً قبل توليد الصور والفيديو.", "kv"),
    ("", "", ""),
    ("تحذير", "", "head"),
    ("ترتيب الأعمدة مش مهم", "عقد Google Sheets بتطابق بالاسم مش بالموضع. بس الاسم لازم يكون حرفياً — "
     "مسافة زيادة أو حرف كبير غلط بيخلي العقدة تعمل عمود جديد بدل ما تكتب في القديم.", "kv"),
    ("متمسحش عمود موجود", "الأعمدة اللي المحرك بيكتبها ومحدش بيقراها (زي warnings و render_url) سيبها — "
     "حذفها بيخلي عقدة الكتابة تفشل.", "kv"),
    ("", "", ""),
    ("لو رفعت الملف ده على Drive كشيت جديد", "", "head"),
    ("هيبقى له Document ID مختلف", "الـ workflow متربط بـ final_engine على "
     "1AdTeb6DaeZPaBpY6JQ3Y5A9SGZrfueBPqm1BiuVTG9w. لو الشيت الجديد هو اللي هيشتغل عليه، "
     "لازم تغيّر documentId في كل عقد Google Sheets (٣٠ عقدة).", "kv"),
    ("الأسهل", "خلّي final_engine زي ما هو، وضيف عليه الأعمدة والتبويبات من هنا. "
     "ساعتها متغيّرش أي حاجة في الـ workflow.", "kv"),
    ("امسح التبويبين دول", "اقرأ-أولاً و قاموس-الأعمدة توثيق مش داتا — امسحهم بعد الرفع، "
     "وامسح صف المثال من كل تبويب.", "kv"),
]
r = 1
for a, b, kind in rows:
    if kind == "title":
        c = ws.cell(row=r, column=1, value=a)
        c.font = Font(name=F, bold=True, size=16, color=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 26
    elif kind == "head":
        c = ws.cell(row=r, column=1, value=a)
        c.font = Font(name=F, bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    elif kind == "kv":
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = Font(name=F, bold=True, size=10, color=NAVY)
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        ca.fill = PatternFill("solid", fgColor=LIGHT)
        cb = ws.cell(row=r, column=2, value=b)
        cb.font = Font(name=F, size=10)
        cb.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 34
    r += 1

# swatches
for row_i, colr in [(6, TEAL), (7, GOLD), (8, NAVY)]:
    ws.cell(row=row_i, column=1).fill = PatternFill("solid", fgColor=colr)
    ws.cell(row=row_i, column=1).font = Font(name=F, bold=True, size=10, color="FFFFFF")

# ---- 2. one sheet per tab ----
ORDER = ["Brand_KB", "Content", "Leads", "Reply_Approval", "Events_Log",
         "Events_Log_Archive", "Handoff", "Performance"]
for tabname in ORDER:
    cols = COLS[tabname]
    ws = wb.create_sheet(tabname)
    style_header(ws, cols)
    example_row(ws, EXAMPLES[tabname], len(cols))

# ---- 3. dictionary ----
ws = wb.create_sheet("قاموس-الأعمدة")
ws.sheet_view.rightToLeft = True
heads = ["التبويب", "العمود", "جديد؟", "مين بيكتبه", "الوصف"]
widths = [20, 26, 10, 16, 74]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(name=F, bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 24

r = 2
for tabname in ORDER:
    for name, owner, is_new, desc in COLS[tabname]:
        vals = [tabname, name, "جديد" if is_new else "", "أنت" if owner == "human" else "المحرك", desc]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=F, size=10, bold=(i == 2))
            c.alignment = Alignment(vertical="top", wrap_text=(i == 5))
            c.border = BORDER
        if is_new:
            ws.cell(row=r, column=3).font = Font(name=F, size=10, bold=True, color="FFFFFF")
            ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=TEAL)
        if owner == "human":
            ws.cell(row=r, column=4).font = Font(name=F, size=10, bold=True, color=NAVY)
            ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FBEEC8")
        r += 1

wb.save(OUT)
n_new = sum(1 for t in ORDER for _, _, is_new, _ in COLS[t] if is_new)
print(f"wrote {OUT} | tabs: {len(ORDER)} | columns: {r-2} | new: {n_new}")
